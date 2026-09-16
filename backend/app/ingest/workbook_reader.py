"""Read the input workbook. Read only, never writes, never repairs a value.

The sheets have a title row, a note row and a blank row before the real headers,
so headers are found by name instead of by a fixed row number.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Sequence

from openpyxl import load_workbook


class WorkbookReadError(Exception):
    """The file could not be opened as an .xlsx workbook."""


@dataclass(frozen=True)
class RawRow:
    """One data row. `number` is the Excel row number, so error messages can point at it."""

    number: int
    values: dict[str, Any]

    def get(self, header: str) -> Any:
        return self.values.get(header)


@dataclass(frozen=True)
class RawTable:
    """A header row found by name, plus the data rows below it."""

    sheet: str
    header_row: int | None
    rows: list[RawRow]
    missing_headers: tuple[str, ...]


def _normalise_header(value: Any) -> str:
    return "" if value is None else str(value).strip().lower()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _row_is_empty(row: Sequence[Any]) -> bool:
    return all(_is_blank(cell) for cell in row)


def read_sheets(source: str | Path | BinaryIO) -> dict[str, list[list[Any]]]:
    """Load every sheet as a list of rows of raw values. The workbook is opened read only."""
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many different types here
        raise WorkbookReadError(str(exc)) from exc

    try:
        sheets: dict[str, list[list[Any]]] = {}
        for name in workbook.sheetnames:
            sheets[name] = [list(row) for row in workbook[name].iter_rows(values_only=True)]
        return sheets
    finally:
        workbook.close()


def find_table(
    sheet: str,
    rows: Iterable[Sequence[Any]],
    required_headers: Sequence[str],
    search_from_row: int = 1,
) -> RawTable:
    """Find the row that carries the required headers, then read data rows below it.

    The best candidate is the row matching the most required headers. Data reading
    stops at the first fully empty row, which is how the two tables of the Station
    sheet are kept apart.
    """
    all_rows = [list(row) for row in rows]
    best_index: int | None = None
    best_matches = 0
    best_positions: dict[str, int] = {}

    for index, row in enumerate(all_rows):
        if index + 1 < search_from_row:
            continue
        positions: dict[str, int] = {}
        for column, cell in enumerate(row):
            header = _normalise_header(cell)
            if header in required_headers and header not in positions:
                positions[header] = column
        if len(positions) > best_matches:
            best_matches = len(positions)
            best_index = index
            best_positions = positions

    if best_index is None or best_matches == 0:
        return RawTable(sheet=sheet, header_row=None, rows=[], missing_headers=tuple(required_headers))

    missing = tuple(header for header in required_headers if header not in best_positions)

    data: list[RawRow] = []
    for index in range(best_index + 1, len(all_rows)):
        row = all_rows[index]
        if _row_is_empty(row):
            break
        values = {
            header: (row[column] if column < len(row) else None)
            for header, column in best_positions.items()
        }
        data.append(RawRow(number=index + 1, values=values))

    return RawTable(
        sheet=sheet,
        header_row=best_index + 1,
        rows=data,
        missing_headers=missing,
    )
